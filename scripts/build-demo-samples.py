# -*- coding: utf-8 -*-
"""Generate only synthetic demo files, or verify their checked-in SHA-256 list.

Generation needs Pillow, python-docx and openpyxl. --check uses only the stdlib.
No user documents, application settings, credentials or network services are read.
"""

import argparse
import csv
import hashlib
import io
import json
import sys
import zipfile
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "samples"
MARKER = "MTO_TEST_ONLY"
NOTICE = "虚构测试资料，仅供软件演示，不用于真实业务"
STAMP = datetime(2026, 10, 15)
PROJECT_A = "workspace/01_项目甲_配套齐全"
PROJECT_B = "workspace/02_项目乙_缺少配套"
OCR_DIR = "workspace/03_待命名单据"
TEXT_DIR = "workspace/04_文本资料"
CASES = [
    {
        "title": "产品采购合同", "profile": "采购合同", "id": "mto-demo-purchase",
        "file": f"{OCR_DIR}/扫描件_001.pdf",
        "lines": ["合同编号：HT-2026-0001", "需方（买方）：测试客户甲有限公司",
                  "供方（卖方）：测试供方有限公司", "签订日期：2026年10月15日"],
        "rows": [["产品名称", "数量", "单价"], ["测试键盘", "10", "125.00"]],
        "total": "合同总额：1250.00 元",
        "fields": {"合同编号": "HT-2026-0001", "需方名称": "测试客户甲有限公司",
                   "签订日期": "20261015", "合同金额": "1250.00"},
        "renamed": "20261015_测试客户甲有限公司_HT-2026-0001.pdf",
    },
    {
        "title": "销售出库送货单", "profile": "销售送货单", "id": "mto-demo-delivery",
        "file": f"{OCR_DIR}/扫描件_002.pdf",
        "lines": ["送货单号：SH-20261016-01", "客户名称：测试客户乙有限公司",
                  "发货单位：测试供方有限公司", "送货日期：2026-10-16"],
        "rows": [["产品名称", "数量", "单位"], ["测试键盘", "2", "件"],
                 ["测试鼠标", "3", "件"]],
        "total": "合计数量：5",
        "fields": {"送货单号": "SH-20261016-01", "客户名称": "测试客户乙有限公司",
                   "送货日期": "20261016", "产品名称": "测试键盘、测试鼠标", "送货数量": "5"},
        "renamed": "20261016_测试客户乙有限公司_SH-20261016-01.pdf",
    },
    {
        "title": "商务设备报价单", "profile": "商务报价单", "id": "mto-demo-quotation",
        "file": f"{OCR_DIR}/扫描件_003.pdf",
        "lines": ["报价单号：BJ-20261017-01", "客户单位：测试客户丙有限公司",
                  "报价单位：测试供方有限公司", "报价日期：2026-10-17",
                  "有效期至：2026-10-31"],
        "rows": [["产品名称", "数量", "单价"], ["测试键盘", "2", "100.00"],
                 ["测试鼠标", "3", "160.00"]],
        "total": "价税合计：680.00 元",
        "fields": {"报价单号": "BJ-20261017-01", "客户单位": "测试客户丙有限公司",
                   "报价日期": "20261017", "价税合计": "680.00"},
        "renamed": "20261017_测试客户丙有限公司_BJ-20261017-01.pdf",
    },
]
PDF_COPY = f"{PROJECT_A}/合同.pdf"
PNG_COPY = f"{OCR_DIR}/扫描件_004.png"
ASSETS = [c["file"] for c in CASES] + [
    PDF_COPY, PNG_COPY, f"{PROJECT_A}/合同.docx", f"{PROJECT_B}/合同.docx",
    f"{PROJECT_A}/报价明细.xlsx", f"{TEXT_DIR}/项目说明.md",
    f"{TEXT_DIR}/物料清单.csv", f"{TEXT_DIR}/使用说明.txt", "ocr-profiles.json",
]


def destination(relative):
    path = OUT / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def json_write(relative, value):
    destination(relative).write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def normalize_zip(path):
    with zipfile.ZipFile(path) as source:
        entries = [(name, source.read(name)) for name in sorted(source.namelist())]
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for name, data in entries:
            entry = zipfile.ZipInfo(name, (2026, 10, 15, 0, 0, 0))
            target.writestr(entry, data, compress_type=zipfile.ZIP_DEFLATED)


def check():
    manifest = json.loads((OUT / "expected.json").read_text(encoding="utf-8"))
    assert manifest["origin"] == MARKER
    assert {item["path"] for item in manifest["files"]} == set(ASSETS)
    for item in manifest["files"]:
        data = (OUT / item["path"]).read_bytes()
        assert hashlib.sha256(data).hexdigest() == item["sha256"], item["path"]
        assert len(data) == item["bytes"], item["path"]
    expected = set(ASSETS) | {"README.md", "expected.json"}
    actual = {p.relative_to(OUT).as_posix() for p in OUT.rglob("*") if p.is_file()}
    assert actual == expected, f"Unexpected/missing sample files: {actual ^ expected}"
    workspace = OUT / "workspace"
    sources = list(workspace.rglob(manifest["missingCompanion"]["sourceName"]))
    missing = sorted(p.parent.name for p in sources
                     if not (p.parent / manifest["missingCompanion"]["companionName"]).is_file())
    assert missing == manifest["missingCompanion"]["directories"], missing
    print(f"PASS: {len(ASSETS) - 1} demo documents + OCR profiles; hashes and file list match.")


def generate(rebuild):
    if any((OUT / path).exists() for path in ASSETS + ["expected.json"]):
        if not rebuild:
            raise SystemExit("Samples exist. Use --check, or --rebuild after checking local changes.")
        check()  # Refuse to overwrite modified samples or unrecognized files.

    from PIL import Image, ImageDraw, ImageFont
    from docx import Document
    from docx.shared import Pt
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    # Read a system font, never a user document. Raster PDFs intentionally have no text layer.
    font_path = Path("C:/Windows/Fonts/msyh.ttc")
    if not font_path.is_file():
        raise SystemExit("Generation requires the Windows Microsoft YaHei font.")
    font = ImageFont.truetype(str(font_path), 38)
    small = ImageFont.truetype(str(font_path), 27)
    title_font = ImageFont.truetype(str(font_path), 52)
    profiles = []
    for case in CASES:
        profile = json.loads((ROOT / "ocr-samples" / f'{case["profile"]}_OCR配置.json')
                             .read_text(encoding="utf-8"))
        profile.update(id=case["id"], name=f'测试样例-{case["profile"]}', pages="1", dpi=200)
        profiles.append(profile)
        page = Image.new("RGB", (1654, 2339), "white")
        draw = ImageDraw.Draw(page)
        draw.text((110, 90), f"{MARKER} | {NOTICE}", font=small, fill="#666666")
        draw.text((110, 175), case["title"], font=title_font, fill="black")
        y = 320
        for text in case["lines"]:
            draw.text((110, y), text, font=font, fill="black")
            y += 100
        y += 65
        for row in case["rows"]:
            for x, text in zip((110, 920, 1210), row):
                draw.text((x, y), text, font=font, fill="black")
            y += 105
        draw.text((110, y + 50), case["total"], font=font, fill="black")
        draw.text((110, 2130), NOTICE, font=small, fill="#666666")
        pdf = io.BytesIO()
        page.save(pdf, format="PDF", resolution=200, quality=92,
                  title=f'{MARKER} {case["title"]}', author="MTO Synthetic Samples",
                  creationDate="D:20261015000000Z", modDate="D:20261015000000Z")
        destination(case["file"]).write_bytes(pdf.getvalue())
        if case is CASES[0]:
            destination(PDF_COPY).write_bytes(pdf.getvalue())
        if case is CASES[1]:
            page.save(destination(PNG_COPY), format="PNG", optimize=True)
        page.close()
    json_write("ocr-profiles.json", profiles)

    for folder, project in [(PROJECT_A, "项目甲"), (PROJECT_B, "项目乙")]:
        doc = Document()
        doc.core_properties.author = doc.core_properties.last_modified_by = "MTO Synthetic Samples"
        doc.core_properties.title = f"{MARKER} {project}合同"
        doc.core_properties.created = doc.core_properties.modified = STAMP
        style = doc.styles["Normal"]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(12)
        doc.add_heading(f"{project}合同（虚构测试）", 0)
        for text in [MARKER, NOTICE, f"项目名称：{project}", *CASES[0]["lines"],
                     "测试用途：核对同目录是否存在合同.pdf。",
                     "合同总额：1250.00 元"]:
            doc.add_paragraph(text)
        path = destination(f"{folder}/合同.docx")
        doc.save(path)
        normalize_zip(path)

    book = Workbook()
    book.properties.creator = book.properties.lastModifiedBy = "MTO Synthetic Samples"
    book.properties.title = f"{MARKER} 报价明细"
    book.properties.created = book.properties.modified = STAMP
    sheet = book.active
    sheet.title = "报价明细"
    for row in [[MARKER, NOTICE], ["产品名称", "数量", "单价", "金额"],
                ["测试键盘", 2, 100, 200], ["测试鼠标", 3, 160, 480], ["合计", 5, "", 680]]:
        sheet.append(row)
    for cell in sheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="355E3B")
    for column in "ABCD":
        sheet.column_dimensions[column].width = 24
    for row in sheet.iter_rows(min_row=3, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = "0.00"
    sheet.freeze_panes = "A3"
    book.create_sheet("使用说明").append([MARKER, NOTICE])
    path = destination(f"{PROJECT_A}/报价明细.xlsx")
    book.save(path)
    normalize_zip(path)

    destination(f"{TEXT_DIR}/项目说明.md").write_text(
        f"# 虚构测试项目\n\n{MARKER}\n\n{NOTICE}。\n\n"
        "- 项目甲：合同.docx 与合同.pdf 配套齐全。\n"
        "- 项目乙：只有合同.docx，故意缺少合同.pdf。\n", encoding="utf-8")
    destination(f"{TEXT_DIR}/使用说明.txt").write_text(
        f"{MARKER}\n{NOTICE}\n本目录所有文档均由固定的虚构内容生成。\n", encoding="utf-8")
    with destination(f"{TEXT_DIR}/物料清单.csv").open("w", encoding="utf-8-sig", newline="") as out:
        csv.writer(out).writerows([["标记", "物料名称", "数量", "说明"],
                                  [MARKER, "测试键盘", 2, NOTICE],
                                  [MARKER, "测试鼠标", 3, NOTICE]])

    ocr_cases = [{"path": c["file"], "profileId": c["id"], "fields": c["fields"],
                  "renamed": c["renamed"]} for c in CASES]
    ocr_cases.append(dict(ocr_cases[1], path=PNG_COPY,
                         renamed=CASES[1]["renamed"].removesuffix(".pdf") + ".png"))
    previews = [
        {"path": f"{PROJECT_A}/合同.docx", "kind": "text", "contains": [MARKER, "项目甲"]},
        {"path": f"{PROJECT_B}/合同.docx", "kind": "text", "contains": [MARKER, "项目乙"]},
        {"path": PDF_COPY, "kind": "image", "pageCount": 1},
        {"path": PNG_COPY, "kind": "image", "pageCount": 1},
        {"path": f"{PROJECT_A}/报价明细.xlsx", "kind": "table",
         "contains": [MARKER, "测试键盘", "测试鼠标", "680"], "sheets": ["报价明细", "使用说明"]},
        {"path": f"{PROJECT_A}/报价明细.xlsx", "sheet": 1, "kind": "table", "contains": [MARKER]},
        {"path": f"{TEXT_DIR}/项目说明.md", "kind": "text", "contains": [MARKER, "项目乙"]},
        {"path": f"{TEXT_DIR}/物料清单.csv", "kind": "table", "contains": [MARKER, "测试键盘", "测试鼠标"]},
        {"path": f"{TEXT_DIR}/使用说明.txt", "kind": "text", "contains": [MARKER, NOTICE]},
    ]
    files = [{"path": p, "bytes": (OUT / p).stat().st_size,
              "sha256": hashlib.sha256((OUT / p).read_bytes()).hexdigest()} for p in ASSETS]
    json_write("expected.json", {
        "origin": MARKER, "notice": NOTICE, "files": files, "ocrCases": ocr_cases,
        "previewChecks": previews,
        "missingCompanion": {"sourceName": "合同.docx", "companionName": "合同.pdf",
                             "directories": ["02_项目乙_缺少配套"]},
    })
    print(f"Generated {len(ASSETS) - 1} synthetic demo documents and OCR profiles.")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    options = parser.add_mutually_exclusive_group()
    options.add_argument("--check", action="store_true")
    options.add_argument("--rebuild", action="store_true")
    args = parser.parse_args()
    if args.check:
        check()
    else:
        generate(args.rebuild)
